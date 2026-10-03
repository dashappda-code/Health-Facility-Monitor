import re
import math
import io
import tempfile

import altair as alt
import numpy as np
import pandas as pd
import streamlit as st

from chart_helpers import data_labels_enabled

# ============================================================
# CONFIGURATION
# ============================================================

POPULATION_SHEET_ID = "18qfha01Czh10i4PDRUpuumtRVwbQv7Pn09jFxGSbXHg"
POPULATION_GID = "846450963"

POPULATION_CSV_URL = (
    f"https://docs.google.com/spreadsheets/d/"
    f"{POPULATION_SHEET_ID}/export"
    f"?format=csv&gid={POPULATION_GID}"
)

DEFAULT_DENOMINATOR = 100_000

DENOMINATOR_OPTIONS = [
    1_000,
    10_000,
    100_000,
    1_000_000,
]

MIN_VALID_YEAR = 2000
MAX_VALID_YEAR = 2100

DATA_LABEL_FONT_SIZE = 10

COLORS = [
    "#1F77B4", "#FF7F0E", "#2CA02C", "#D62728",
    "#9467BD", "#8C564B", "#E377C2", "#7F7F7F",
    "#BCBD22", "#17BECF", "#4C78A4", "#F58518",
]

MONTH_NAMES = {
    1: "Jan", 2: "Feb", 3: "Mar", 4: "Apr",
    5: "May", 6: "Jun", 7: "Jul", 8: "Aug",
    9: "Sep", 10: "Oct", 11: "Nov", 12: "Dec",
}


# ============================================================
# LOCAL UI HELPERS
# ============================================================

def _set_widget_state(widget_key, values):
    st.session_state[widget_key] = list(values)

def _reset_selection(widget_key, applied_key, default_values):
    defaults = list(default_values)
    st.session_state[widget_key] = defaults
    st.session_state[applied_key] = defaults

def _apply_deferred_selection(widget_key, applied_key, select_all_key, options):
    if st.session_state.get(select_all_key):
        st.session_state[widget_key] = list(options)
        st.session_state[applied_key] = list(options)
    else:
        st.session_state[applied_key] = list(st.session_state.get(widget_key, []))

def _reset_compact_selection(widget_key, applied_key, select_all_key, default_values):
    defaults = list(default_values)
    st.session_state[widget_key] = defaults
    st.session_state[applied_key] = defaults
    st.session_state[select_all_key] = False

def _smart_year_slider(label, options, default=None, key=None):
    options = sorted(list(options))
    if not options: return None
    if len(options) == 1:
        st.info(f"{label}: **{options[0]}** (Only one year available)")
        return options[0]
    if default not in options: default = options[-1]
    if key is None: key = f"_smart_slider_{label}"
    widget_key = f"{key}__widget"
    if widget_key not in st.session_state: st.session_state[widget_key] = default
    selected = st.select_slider(label, options=options, value=st.session_state[widget_key], key=widget_key)
    st.caption(f"✅ **Selected Year:** {selected}")
    return selected

def _single_select(label, options, default=None, key=None, reset_button_key=None, reset_label="↺ Reset", show_selected=True):
    options = list(options)
    if not options: return None
    if default not in options: default = options[0]
    if key is None: key = f"_single_select_{label}"
    widget_key = f"{key}__widget"

    applied = st.session_state.get(key)
    if isinstance(applied, (list, tuple)): applied = [value for value in applied if value in options][:1]
    elif applied in options: applied = [applied]
    else: applied = [default]

    if not applied: applied = [default]
    if widget_key not in st.session_state: st.session_state[widget_key] = list(applied)

    if reset_button_key is not None:
        st.button(reset_label, key=reset_button_key, help=f"Reset {label} to the default selection.", on_click=_reset_selection, args=(widget_key, key, (default,)))

    selected = st.multiselect(label, options=options, max_selections=1, select_all=False, placeholder="Select one option", key=widget_key)
    applied_value = selected[0] if selected else None
    if applied_value is not None: st.session_state[key] = [applied_value]

    if show_selected:
        if applied_value is None: st.caption("☐ **Selected:** None")
        else: st.caption(f"✅ **Selected:** {applied_value} · Charts below use this selection.")
    return applied_value

def _compact_multiselect(label, options, default=None, key=None, reset_button_key=None, reset_label="↺ Reset", show_selected=True, defer=False, apply_label="Apply Selection"):
    options = list(options)
    if not options: return []
    if default is None: default = options[: min(5, len(options))]
    default = [value for value in default if value in options]
    if key is None: key = f"_compact_multiselect_{label}"

    widget_key, select_all_key, applied_key = f"{key}__widget", f"{key}__select_all", key
    applied = st.session_state.get(applied_key)
    if isinstance(applied, (list, tuple)): applied = [value for value in applied if value in options]
    else: applied = list(default)

    if widget_key not in st.session_state: st.session_state[widget_key] = list(applied)

    if defer:
        form_key = f"{key}__form"
        with st.form(form_key, clear_on_submit=False):
            st.checkbox("☑ Select All Options", key=select_all_key)
            selected = st.multiselect(label, options=options, key=widget_key)
            apply_col, reset_col = st.columns(2)
            with apply_col: st.form_submit_button(f"✓ {apply_label}", use_container_width=True, on_click=_apply_deferred_selection, args=(widget_key, applied_key, select_all_key, options))
            with reset_col: st.form_submit_button(reset_label, use_container_width=True, on_click=_reset_compact_selection, args=(widget_key, applied_key, select_all_key, tuple(default)))
        selected = st.session_state.get(applied_key, list(default))
    else:
        def _toggle_all():
            if st.session_state.get(select_all_key): st.session_state[widget_key] = list(options)
            else: st.session_state[widget_key] = []
                
        col1, col2 = st.columns([1, 1])
        with col1: st.checkbox("☑ Select All Options", key=select_all_key, on_change=_toggle_all)
        with col2:
            if reset_button_key is not None: st.button(reset_label, key=reset_button_key, use_container_width=True, on_click=_reset_compact_selection, args=(widget_key, applied_key, select_all_key, tuple(default)))
        selected = st.multiselect(label, options=options, key=widget_key)
        st.session_state[applied_key] = list(selected)

    selected = [value for value in selected if value in options]
    if show_selected:
        if selected:
            display = ", ".join(str(value) for value in selected)
            if len(selected) > 5: display = f"{len(selected)} options selected"
            st.caption(f"✅ **Selected:** {display} · The chart uses all checked selections.")
        else: st.caption("☐ **Selected:** None")
    return selected

def _chart_title(title, x_title, y_title, interpretation=None):
    subtitle_parts = [f"X-axis: {x_title}", f"Y-axis: {y_title}"]
    if interpretation: subtitle_parts.append(f"Interpretation: {interpretation}")
    return alt.TitleParams(text=title, subtitle=" | ".join(subtitle_parts), anchor="start", fontSize=16, fontWeight="bold", subtitleFontSize=11, subtitleFontWeight="normal", subtitlePadding=6, limit=1000)

# ============================================================
# BASIC HELPERS
# ============================================================

def _clean_series(series):
    if series is None: return pd.Series(dtype="object")
    return series.fillna("").astype(str).str.strip()

def _valid_text(series):
    values = _clean_series(series)
    return values.ne("") & values.ne("nan") & values.ne("NaT") & values.ne("None") & values.ne("-")

def _safe_numeric(series):
    if series is None: return pd.Series(dtype=float)
    cleaned = series.astype(str).str.replace(",", "", regex=False).str.replace(" ", "", regex=False).str.strip().replace({"": np.nan, "-": np.nan, "nan": np.nan, "NaN": np.nan, "None": np.nan, "NA": np.nan, "N/A": np.nan})
    return pd.to_numeric(cleaned, errors="coerce")

def _alphabetical(values):
    cleaned = []
    for value in values:
        text = str(value).strip()
        if text and text not in {"nan", "NaT", "None", "-"}: cleaned.append(text)
    return sorted(list(dict.fromkeys(cleaned)), key=lambda x: x.upper())

def _month_number(value):
    if pd.isna(value): return np.nan
    text = str(value).strip().lower()
    month_map = {"january": 1, "jan": 1, "1": 1, "01": 1, "february": 2, "feb": 2, "2": 2, "02": 2, "march": 3, "mar": 3, "3": 3, "03": 3, "april": 4, "apr": 4, "4": 4, "04": 4, "may": 5, "5": 5, "05": 5, "june": 6, "jun": 6, "6": 6, "06": 6, "july": 7, "jul": 7, "7": 7, "07": 7, "august": 8, "aug": 8, "8": 8, "08": 8, "september": 9, "sep": 9, "sept": 9, "9": 9, "09": 9, "october": 10, "oct": 10, "10": 10, "november": 11, "nov": 11, "11": 11, "december": 12, "dec": 12, "12": 12}
    if text in month_map: return month_map[text]
    try:
        number = int(float(text))
        if 1 <= number <= 12: return number
    except Exception: pass
    return np.nan

def _format_number(value, decimals=0):
    if pd.isna(value): return "Unavailable"
    try:
        if decimals == 0: return f"{float(value):,.0f}"
        return f"{float(value):,.{decimals}f}"
    except Exception: return str(value)

# ============================================================
# CASE DATA PREPARATION
# ============================================================

@st.cache_data(ttl=1800, show_spinner=False)
def _prepare_case_data(df):
    if df is None or df.empty: return pd.DataFrame()
    temp = df.copy()
    if "Year" in temp.columns: temp["Analysis_Year"] = pd.to_numeric(temp["Year"], errors="coerce")
    elif "Reporting Date" in temp.columns:
        dates = pd.to_datetime(temp["Reporting Date"], errors="coerce", dayfirst=True)
        temp["Analysis_Year"] = dates.dt.year
    else: temp["Analysis_Year"] = np.nan

    if "Month" in temp.columns: temp["Analysis_Month"] = temp["Month"].map(_month_number)
    elif "Reporting Date" in temp.columns:
        dates = pd.to_datetime(temp["Reporting Date"], errors="coerce", dayfirst=True)
        temp["Analysis_Month"] = dates.dt.month
    else: temp["Analysis_Month"] = np.nan

    temp = temp[temp["Analysis_Year"].between(MIN_VALID_YEAR, MAX_VALID_YEAR)].copy()
    if temp.empty: return pd.DataFrame()

    temp["Analysis_Year"] = temp["Analysis_Year"].astype(int)
    if "Ward Name" in temp.columns: temp["Ward Name"] = _clean_series(temp["Ward Name"])
    if "Disease" in temp.columns: temp["Disease"] = _clean_series(temp["Disease"])
    if "Facility Name" in temp.columns: temp["Facility Name"] = _clean_series(temp["Facility Name"])
    return temp

# ============================================================
# POPULATION DATA
# ============================================================

@st.cache_data(ttl=3600, show_spinner=False)
def _load_population_raw():
    population_df = pd.read_csv(POPULATION_CSV_URL)
    population_df.columns = [str(column).strip() for column in population_df.columns]
    return population_df

def _find_ward_column(df):
    if df is None or df.empty: return None
    for candidate in ["WARD", "Ward", "Ward Name", "WARD NAME", "ward"]:
        if candidate in df.columns: return candidate
    for column in df.columns:
        if str(column).strip().lower().replace("_", " ") in {"ward", "ward name"}: return column
    return None

def _extract_population_year(column_name):
    match = re.search(r"(20\d{2})", str(column_name))
    if not match: return None
    year = int(match.group(1))
    if MIN_VALID_YEAR <= year <= MAX_VALID_YEAR: return year
    return None

def _prepare_population_data(population_df):
    if population_df is None or population_df.empty: return pd.DataFrame(), pd.DataFrame(), []
    source = population_df.copy()
    ward_column = _find_ward_column(source)
    if ward_column is None: return pd.DataFrame(), pd.DataFrame(), []
    
    year_columns = {}
    for column in source.columns:
        if column == ward_column: continue
        year = _extract_population_year(column)
        if year is not None and year not in year_columns: year_columns[year] = column
            
    if not year_columns: return pd.DataFrame(), pd.DataFrame(), []
    
    source["WARD"] = _clean_series(source[ward_column])
    source = source[_valid_text(source["WARD"])].copy()
    source = source[~source["WARD"].str.upper().isin({"TOTAL", "GRAND TOTAL", "MUMBAI TOTAL", "ALL"})].copy()
    available_years = sorted(year_columns.keys())
    
    wide_data = {"WARD": source["WARD"]}
    for year in available_years: wide_data[str(year)] = _safe_numeric(source[year_columns[year]])
        
    population_wide = pd.DataFrame(wide_data)
    numeric_columns = [str(year) for year in available_years]
    population_wide = population_wide.groupby("WARD", as_index=False)[numeric_columns].sum(min_count=1)
    
    population_long = population_wide.melt(id_vars=["WARD"], value_vars=numeric_columns, var_name="Population_Year", value_name="Population")
    population_long["Population_Year"] = pd.to_numeric(population_long["Population_Year"], errors="coerce")
    population_long["Population"] = pd.to_numeric(population_long["Population"], errors="coerce")
    population_long = population_long[population_long["Population_Year"].notna()].copy()
    population_long["Population_Year"] = population_long["Population_Year"].astype(int)
    population_long.loc[population_long["Population"] <= 0, "Population"] = np.nan
    return population_wide, population_long, available_years

# ============================================================
# POPULATION MATCHING
# ============================================================

def _population_for_ward_year(population_long, ward, analysis_year):
    if population_long is None or population_long.empty: return np.nan, None, "Population unavailable"
    ward_text = str(ward).strip()
    try: analysis_year = int(analysis_year)
    except Exception: return np.nan, None, "Invalid analysis year"
    
    data = population_long[_clean_series(population_long["WARD"]) == ward_text].copy()
    data = data[data["Population"].notna() & data["Population"].gt(0)]
    if data.empty: return np.nan, None, "Population unavailable"
    
    exact = data[data["Population_Year"] == analysis_year]
    if not exact.empty:
        row = exact.iloc[0]
        return float(row["Population"]), int(row["Population_Year"]), "Exact year"
        
    previous = data[data["Population_Year"] < analysis_year]
    if not previous.empty:
        selected_year = int(previous["Population_Year"].max())
        row = previous[previous["Population_Year"] == selected_year].iloc[0]
        return float(row["Population"]), selected_year, "Previous-year fallback"
        
    future = data[data["Population_Year"] > analysis_year]
    if not future.empty:
        selected_year = int(future["Population_Year"].min())
        row = future[future["Population_Year"] == selected_year].iloc[0]
        return float(row["Population"]), selected_year, "Future-year fallback"
        
    return np.nan, None, "Population unavailable"

# ============================================================
# INCIDENCE TABLES
# ============================================================

def _build_ward_year_incidence(case_df, population_long, denominator):
    if case_df is None or case_df.empty or "Ward Name" not in case_df.columns: return pd.DataFrame()
    temp = case_df[_valid_text(case_df["Ward Name"])].copy()
    if temp.empty: return pd.DataFrame()
    grouped = temp.groupby(["Ward Name", "Analysis_Year"], observed=True).size().reset_index(name="Cases")
    rows = []
    for _, row in grouped.iterrows():
        population, year_used, status = _population_for_ward_year(population_long, row["Ward Name"], row["Analysis_Year"])
        rate = np.nan
        if pd.notna(population) and population > 0: rate = row["Cases"] / population * denominator
        rows.append({"Ward": row["Ward Name"], "Year": int(row["Analysis_Year"]), "Cases": int(row["Cases"]), "Population": population, "Population Year Used": year_used, "Population Match Status": status, "Incidence Rate": (round(rate, 2) if pd.notna(rate) else np.nan)})
    return pd.DataFrame(rows)

def _build_disease_ward_year_incidence(case_df, population_long, denominator):
    if case_df is None or case_df.empty or not {"Ward Name", "Disease", "Analysis_Year"}.issubset(case_df.columns): return pd.DataFrame()
    temp = case_df[_valid_text(case_df["Ward Name"]) & _valid_text(case_df["Disease"])].copy()
    if temp.empty: return pd.DataFrame()
    grouped = temp.groupby(["Ward Name", "Disease", "Analysis_Year"], observed=True).size().reset_index(name="Cases")
    rows = []
    for _, row in grouped.iterrows():
        population, year_used, status = _population_for_ward_year(population_long, row["Ward Name"], row["Analysis_Year"])
        rate = np.nan
        if pd.notna(population) and population > 0: rate = row["Cases"] / population * denominator
        rows.append({"Ward": row["Ward Name"], "Disease": row["Disease"], "Year": int(row["Analysis_Year"]), "Cases": int(row["Cases"]), "Population": population, "Population Year Used": year_used, "Population Match Status": status, "Incidence Rate": (round(rate, 2) if pd.notna(rate) else np.nan)})
    return pd.DataFrame(rows)

def _build_monthly_disease_ward_incidence(case_df, population_long, denominator):
    if case_df is None or case_df.empty or not {"Ward Name", "Disease", "Analysis_Year", "Analysis_Month"}.issubset(case_df.columns): return pd.DataFrame()
    temp = case_df[case_df["Analysis_Month"].between(1, 12) & _valid_text(case_df["Ward Name"]) & _valid_text(case_df["Disease"])].copy()
    if temp.empty: return pd.DataFrame()
    grouped = temp.groupby(["Ward Name", "Disease", "Analysis_Year", "Analysis_Month"], observed=True).size().reset_index(name="Cases")
    rows = []
    for _, row in grouped.iterrows():
        population, year_used, status = _population_for_ward_year(population_long, row["Ward Name"], row["Analysis_Year"])
        rate = np.nan
        if pd.notna(population) and population > 0: rate = row["Cases"] / population * denominator
        year, month = int(row["Analysis_Year"]), int(row["Analysis_Month"])
        date_value = pd.Timestamp(year, month, 1)
        rows.append({"Ward": row["Ward Name"], "Disease": row["Disease"], "Year": year, "Month": month, "Date": date_value, "Period": date_value.strftime("%b-%Y"), "Cases": int(row["Cases"]), "Population": population, "Population Year Used": year_used, "Population Match Status": status, "Incidence Rate": (round(rate, 2) if pd.notna(rate) else np.nan)})
    return pd.DataFrame(rows)

def _build_monthly_ward_incidence(case_df, population_long, denominator):
    if case_df is None or case_df.empty or not {"Ward Name", "Analysis_Year", "Analysis_Month"}.issubset(case_df.columns): return pd.DataFrame()
    temp = case_df[case_df["Analysis_Month"].between(1, 12) & _valid_text(case_df["Ward Name"])].copy()
    if temp.empty: return pd.DataFrame()
    grouped = temp.groupby(["Ward Name", "Analysis_Year", "Analysis_Month"], observed=True).size().reset_index(name="Cases")
    rows = []
    for _, row in grouped.iterrows():
        population, year_used, status = _population_for_ward_year(population_long, row["Ward Name"], row["Analysis_Year"])
        rate = np.nan
        if pd.notna(population) and population > 0: rate = row["Cases"] / population * denominator
        year, month = int(row["Analysis_Year"]), int(row["Analysis_Month"])
        date_value = pd.Timestamp(year, month, 1)
        rows.append({"Ward": row["Ward Name"], "Year": year, "Month": month, "Date": date_value, "Period": date_value.strftime("%b-%Y"), "Cases": int(row["Cases"]), "Population": population, "Population Year Used": year_used, "Population Match Status": status, "Incidence Rate": (round(rate, 2) if pd.notna(rate) else np.nan)})
    result = pd.DataFrame(rows)
    if result.empty: return result
    return result.sort_values(["Date", "Ward"]).reset_index(drop=True)

def _aggregate_monthly_incidence(monthly_source, denominator):
    if monthly_source is None or monthly_source.empty or not {"Date", "Period", "Cases", "Population"}.issubset(monthly_source.columns): return pd.DataFrame()
    valid = monthly_source[monthly_source["Population"].notna() & monthly_source["Population"].gt(0)].copy()
    if valid.empty: return pd.DataFrame()
    result = valid.groupby(["Date", "Period"], as_index=False).agg(Cases=("Cases", "sum"), Population=("Population", "sum")).sort_values("Date").reset_index(drop=True)
    result["Cases"] = result["Cases"].astype(int)
    result["Incidence Rate"] = (result["Cases"] / result["Population"] * denominator).round(2)
    return result

def _aggregate_year_incidence(ward_incidence, denominator):
    if ward_incidence is None or ward_incidence.empty: return pd.DataFrame()
    rows = []
    for year, group in ward_incidence.groupby("Year", observed=True):
        valid = group[group["Population"].notna() & group["Population"].gt(0)]
        if valid.empty: continue
        cases, population = int(valid["Cases"].sum()), float(valid["Population"].sum())
        rows.append({"Year": int(year), "Cases": cases, "Population": population, "Incidence Rate": round(cases / population * denominator, 2)})
    if not rows: return pd.DataFrame()
    return pd.DataFrame(rows).sort_values("Year").reset_index(drop=True)

@st.cache_data(ttl=1800, show_spinner=False)
def _build_incidence_bundle(case_df, population_long, denominator):
    return (
        _build_ward_year_incidence(case_df, population_long, denominator),
        _build_disease_ward_year_incidence(case_df, population_long, denominator),
        _build_monthly_ward_incidence(case_df, population_long, denominator),
        _build_monthly_disease_ward_incidence(case_df, population_long, denominator),
        _aggregate_year_incidence(_build_ward_year_incidence(case_df, population_long, denominator), denominator),
        _population_audit(case_df, population_long)
    )

# ============================================================
# CHART CONFIG & RENDERING
# ============================================================

def _legend_columns(series_count):
    return 1 if max(1, int(series_count or 1)) == 1 else max(2, math.ceil(series_count / 2))

def _legend(series_count, title=None):
    return alt.Legend(title=title, orient="bottom", direction="horizontal", columns=_legend_columns(series_count), labelFontSize=9, titleFontSize=10, symbolSize=65, labelLimit=180, columnPadding=12, rowPadding=6, padding=8)

def _finalize_chart(chart, series_count=1):
    return chart.configure_legend(orient="bottom", direction="horizontal", columns=_legend_columns(max(1, series_count)), labelFontSize=9, titleFontSize=10, symbolSize=65, labelLimit=180, columnPadding=12, rowPadding=6, padding=8, offset=8).configure_title(anchor="start", fontSize=16, subtitleFontSize=11, subtitlePadding=6).configure_view(stroke=None)

def _render_bar_chart(dataframe, x_column, y_column, title, y_title, height=430):
    if dataframe is None or dataframe.empty or x_column not in dataframe.columns or y_column not in dataframe.columns: return
    chart_df = dataframe[dataframe[y_column].notna()].copy()
    if chart_df.empty: return
    chart_df[x_column] = chart_df[x_column].astype(str)
    order = chart_df[x_column].drop_duplicates().tolist()
    scale = alt.Scale(domain=order, range=[COLORS[i % len(COLORS)] for i in range(len(order))])
    
    bars = alt.Chart(chart_df).mark_bar().encode(
        x=alt.X(f"{x_column}:N", sort=order, axis=alt.Axis(title=x_column, labelAngle=-45, labelLimit=180)),
        y=alt.Y(f"{y_column}:Q", axis=alt.Axis(title=y_title)),
        color=alt.Color(f"{x_column}:N", scale=scale, legend=_legend(len(order), x_column)),
        tooltip=[alt.Tooltip(f"{x_column}:N", title=x_column), alt.Tooltip(f"{y_column}:Q", title=y_title, format=",.2f")],
    ).properties(height=height, title=_chart_title(title, x_column, y_title))
    
    chart = bars + alt.Chart(chart_df).mark_text(dy=-8, fontSize=DATA_LABEL_FONT_SIZE, fontWeight="bold").encode(x=alt.X(f"{x_column}:N", sort=order), y=alt.Y(f"{y_column}:Q"), text=alt.Text(f"{y_column}:Q", format=",.2f"), color=alt.Color(f"{x_column}:N", scale=scale, legend=None)) if data_labels_enabled() else bars
    chart = _finalize_chart(chart, series_count=len(order))

    if "export_charts" not in st.session_state: st.session_state["export_charts"] = {}
    st.session_state["export_charts"][title] = chart
    st.altair_chart(chart, use_container_width=True)

def _render_line_chart(dataframe, x_column, series_column, value_column, title, y_title, height=450):
    if dataframe is None or dataframe.empty or not {x_column, series_column, value_column}.issubset(dataframe.columns): return
    chart_df = dataframe[dataframe[value_column].notna()].copy()
    if chart_df.empty: return
    series_order = chart_df[series_column].astype(str).drop_duplicates().tolist()
    scale = alt.Scale(domain=series_order, range=[COLORS[i % len(COLORS)] for i in range(len(series_order))])
    
    lines = alt.Chart(chart_df).mark_line(point=True, strokeWidth=3).encode(
        x=alt.X(f"{x_column}:N", sort=None, axis=alt.Axis(title=x_column, labelAngle=-45, labelLimit=150)),
        y=alt.Y(f"{value_column}:Q", axis=alt.Axis(title=y_title)),
        color=alt.Color(f"{series_column}:N", sort=series_order, scale=scale, legend=_legend(len(series_order), series_column)),
        tooltip=[alt.Tooltip(f"{x_column}:N", title=x_column), alt.Tooltip(f"{series_column}:N", title=series_column), alt.Tooltip(f"{value_column}:Q", title=y_title, format=",.2f")],
    ).properties(height=height, title=_chart_title(title, x_column, y_title))
    
    chart = lines + alt.Chart(chart_df).mark_text(dy=-9, fontSize=DATA_LABEL_FONT_SIZE, fontWeight="bold").encode(x=alt.X(f"{x_column}:N", sort=None), y=alt.Y(f"{value_column}:Q"), text=alt.Text(f"{value_column}:Q", format=",.2f"), color=alt.Color(f"{series_column}:N", scale=scale, legend=None)) if data_labels_enabled() else lines
    chart = _finalize_chart(chart, series_count=len(series_order))

    if "export_charts" not in st.session_state: st.session_state["export_charts"] = {}
    st.session_state["export_charts"][title] = chart
    st.altair_chart(chart, use_container_width=True)

# ============================================================
# OBSERVATIONS & METHODOLOGY
# ============================================================

def _observation_box(text):
    if text: st.info(f"📌 **Chart Observation:** {text}")

def _bar_observation(dataframe, category, value, unit):
    if dataframe is None or dataframe.empty or category not in dataframe.columns or value not in dataframe.columns: return None
    valid = dataframe[dataframe[value].notna()].copy()
    if valid.empty: return None
    highest, lowest = valid.loc[valid[value].idxmax()], valid.loc[valid[value].idxmin()]
    if len(valid) == 1: return f"{highest[category]} recorded {_format_number(highest[value], 2)} {unit}."
    return f"The highest value is observed for {highest[category]} ({_format_number(highest[value], 2)} {unit}), while the lowest among the displayed categories is {lowest[category]} ({_format_number(lowest[value], 2)} {unit})."

def _trend_observation(dataframe, period_column, value_column):
    if dataframe is None or dataframe.empty or value_column not in dataframe.columns: return None
    valid = dataframe[dataframe[value_column].notna()].copy()
    if valid.empty: return None
    peak, first, last = valid.loc[valid[value_column].idxmax()], valid.iloc[0], valid.iloc[-1]
    direction = "higher" if float(last[value_column]) > float(first[value_column]) else "lower" if float(last[value_column]) < float(first[value_column]) else "the same"
    return f"The highest displayed value occurs in {peak[period_column]} ({_format_number(peak[value_column], 2)}). The last displayed value is {direction} than the first displayed value."

def _methodology_expander(title, lines):
    with st.expander(f"ℹ️ Methodology — {title}", expanded=False):
        for line in lines: st.markdown(f"- {line}")

# ============================================================
# STATISTICAL FUNCTIONS
# ============================================================

def _apply_baseline(source, baseline):
    if source is None or source.empty: return source
    years = sorted(source["Analysis_Year"].dropna().astype(int).unique().tolist())
    if not years or baseline == "All Available Years": return source.copy()
    count = {"Last 2 Years": 2, "Last 3 Years": 3, "Last 5 Years": 5}.get(baseline)
    return source[source["Analysis_Year"].isin(years[-count:])].copy() if count else source.copy()

def _monthly_case_series(source):
    if source is None or source.empty or not {"Analysis_Year", "Analysis_Month"}.issubset(source.columns): return pd.DataFrame()
    temp = source[source["Analysis_Month"].between(1, 12)].copy()
    if temp.empty: return pd.DataFrame()
    monthly = temp.groupby(["Analysis_Year", "Analysis_Month"], observed=True).size().reset_index(name="Cases")
    monthly["Date"] = pd.to_datetime(dict(year=monthly["Analysis_Year"].astype(int), month=monthly["Analysis_Month"].astype(int), day=1))
    complete = pd.DataFrame({"Date": pd.date_range(monthly["Date"].min(), monthly["Date"].max(), freq="MS")})
    complete["Analysis_Year"], complete["Analysis_Month"] = complete["Date"].dt.year, complete["Date"].dt.month
    complete = complete.merge(monthly[["Analysis_Year", "Analysis_Month", "Cases"]], how="left", on=["Analysis_Year", "Analysis_Month"])
    complete["Cases"] = complete["Cases"].fillna(0).astype(int)
    complete["Period"] = complete["Date"].dt.strftime("%b-%Y")
    return complete

def _threshold_statistics(monthly):
    if monthly is None or monthly.empty: return None
    values = pd.to_numeric(monthly["Cases"], errors="coerce").dropna()
    if values.empty: return None
    mean, sd = float(values.mean()), float(values.std(ddof=1)) if len(values) > 1 else 0.0
    return {"Mean": mean, "SD": sd, "Mean + 1 SD": mean + sd, "Mean + 2 SD": mean + (2 * sd), "Mean + 3 SD": mean + (3 * sd)}

def _render_threshold_chart(monthly, statistics, title):
    if monthly is None or monthly.empty or statistics is None: return
    chart_df = monthly[["Period", "Cases"]].copy()
    for key in ["Mean", "Mean + 1 SD", "Mean + 2 SD", "Mean + 3 SD"]: chart_df[key] = statistics[key]
    long_df = chart_df.melt(id_vars=["Period"], value_vars=["Cases", "Mean", "Mean + 1 SD", "Mean + 2 SD", "Mean + 3 SD"], var_name="Series", value_name="Value")
    series_order = ["Cases", "Mean", "Mean + 1 SD", "Mean + 2 SD", "Mean + 3 SD"]
    scale = alt.Scale(domain=series_order, range=["#1F77B4", "#2CA02C", "#FFBF00", "#FF7F0E", "#D62728"])
    
    chart = alt.Chart(long_df).mark_line(strokeWidth=3).encode(
        x=alt.X("Period:N", sort=None, axis=alt.Axis(title="Period", labelAngle=-45, labelLimit=140)),
        y=alt.Y("Value:Q", title="Monthly Cases"),
        color=alt.Color("Series:N", sort=series_order, scale=scale, legend=_legend(len(series_order), "Statistical Series")),
        tooltip=[alt.Tooltip("Period:N", title="Period"), alt.Tooltip("Series:N", title="Series"), alt.Tooltip("Value:Q", title="Value", format=",.2f")],
    ).properties(height=460, title=_chart_title(title, "Period", "Monthly Cases", "Observed monthly cases compared with Mean and SD thresholds."))
    
    case_points = long_df[long_df["Series"] == "Cases"]
    chart += alt.Chart(case_points).mark_point(filled=True, size=70).encode(x=alt.X("Period:N", sort=None), y="Value:Q", color=alt.Color("Series:N", scale=scale, legend=None))
    if data_labels_enabled(): chart += alt.Chart(case_points).mark_text(dy=-10, fontSize=DATA_LABEL_FONT_SIZE, fontWeight="bold").encode(x=alt.X("Period:N", sort=None), y=alt.Y("Value:Q"), text=alt.Text("Value:Q", format=",.1f"), color=alt.Color("Series:N", scale=scale, legend=None))
        
    latest_period = chart_df["Period"].iloc[-1]
    latest_labels = pd.DataFrame({"Period": [latest_period] * len(series_order), "Series": series_order, "Value": [statistics.get(s, chart_df.loc[chart_df["Period"] == latest_period, "Cases"].iloc[0]) if s != "Cases" else chart_df.loc[chart_df["Period"] == latest_period, "Cases"].iloc[0] for s in series_order]})
    latest_labels["Label"] = latest_labels.apply(lambda r: (f"Observed = {r['Value']:,.1f}" if r["Series"] == "Cases" else f"{r['Series']} = {r['Value']:,.1f}"), axis=1)
    chart += alt.Chart(latest_labels).mark_text(align="left", dx=8, fontSize=10, fontWeight="bold").encode(x=alt.X("Period:N", sort=None), y=alt.Y("Value:Q"), text=alt.Text("Label:N"), color=alt.Color("Series:N", scale=scale, legend=None))
    
    chart = _finalize_chart(chart, series_count=len(series_order))

    if "export_charts" not in st.session_state: st.session_state["export_charts"] = {}
    st.session_state["export_charts"][title] = chart
    st.caption("How to read this chart: blue line = observed cases; green = Mean; yellow/orange/red = Mean + 1/2/3 SD.")
    st.altair_chart(chart, use_container_width=True)

def _threshold_observation(monthly, statistics):
    if monthly is None or monthly.empty or statistics is None: return None
    latest = monthly.iloc[-1]
    latest_cases = float(latest["Cases"])
    status = "above the Mean + 3 SD statistical level" if latest_cases > statistics["Mean + 3 SD"] else "above the Mean + 2 SD statistical level" if latest_cases > statistics["Mean + 2 SD"] else "above the Mean + 1 SD statistical level" if latest_cases > statistics["Mean + 1 SD"] else "within or below the Mean + 1 SD statistical level"
    peak = monthly.loc[monthly["Cases"].idxmax()]
    return f"The latest period ({latest['Period']}) recorded {int(latest_cases):,} cases and is {status}. The highest monthly count in the displayed baseline was {int(peak['Cases']):,} cases in {peak['Period']}. These are descriptive statistical signals and do not independently confirm an outbreak."

def _population_audit(case_df, population_long):
    if case_df is None or case_df.empty or "Ward Name" not in case_df.columns: return pd.DataFrame()
    combinations = case_df[["Ward Name", "Analysis_Year"]].copy()
    combinations = combinations[_valid_text(combinations["Ward Name"])].drop_duplicates().sort_values(["Analysis_Year", "Ward Name"])
    rows = [{"Ward": row["Ward Name"], "Analysis Year": int(row["Analysis_Year"]), "Population": p, "Population Year Used": y, "Status": s} for _, row in combinations.iterrows() for p, y, s in [_population_for_ward_year(population_long, row["Ward Name"], row["Analysis_Year"])]]
    return pd.DataFrame(rows)

# ============================================================
# MAIN INCIDENCE PAGE
# ============================================================

def render_incidence_analysis(df):
    st.session_state["export_charts"] = {} # Clears old charts on new render
    
    st.subheader("📐 Incidence & Statistical Surveillance")
    st.caption("This section answers three management questions: 1) What is the disease burden relative to population? 2) Where and when is incidence higher? 3) Are recent monthly case counts above the historical descriptive statistical baseline?")

    if df is None or df.empty:
        st.warning("No records are available for the selected filters.")
        return

    case_df = _prepare_case_data(df)
    if case_df.empty:
        st.warning("Valid year information is required for incidence analysis.")
        return

    population_error = None
    try:
        population_raw = _load_population_raw()
        population_wide, population_long, population_years = _prepare_population_data(population_raw)
    except Exception as error:
        population_error, population_wide, population_long, population_years = error, pd.DataFrame(), pd.DataFrame(), []

    # --- 1. OVERVIEW ---
    st.markdown("### 1. 📊 Population & Incidence Overview")
    st.info("Purpose: Incidence relates the number of cases to the population of the corresponding ward. This allows wards of different population sizes to be compared on a common rate scale.")

    denominator_label = _single_select("Incidence Rate Denominator", ["Per 1,000", "Per 10,000", "Per 100,000", "Per 1,000,000"], default="Per 100,000", key="phase8b_denom_state")
    denominator = {"Per 1,000": 1_000, "Per 10,000": 10_000, "Per 100,000": 100_000, "Per 1,000,000": 1_000_000}.get(denominator_label, DEFAULT_DENOMINATOR)

    case_years = sorted(case_df["Analysis_Year"].unique().tolist())
    population_wards = population_long["WARD"].nunique() if not population_long.empty else 0

    c1, c2, c3, c4 = st.columns(4)
    with c1: st.metric("Filtered Records", f"{len(case_df):,}")
    with c2: st.metric("Case Years", f"{len(case_years):,}")
    with c3: st.metric("Population Wards", f"{population_wards:,}")
    with c4: st.metric("Latest Population Year", str(max(population_years)) if population_years else "Unavailable")

    if population_error is not None: st.error("Population data could not be loaded from the Population worksheet.")
    elif population_long.empty: st.warning("No valid population data or population-year columns were detected.")
    else:
        st.success(f"Population dataset loaded successfully. Detected years: {', '.join(str(x) for x in population_years)}")
        with st.expander("📋 View Complete Population Dataset", expanded=False): st.dataframe(population_wide, use_container_width=True, hide_index=True)

    _methodology_expander("Population & Incidence Overview", ["Incidence Rate = Cases ÷ Population × selected denominator.", "The default denominator is 100,000 population.", "Population year columns are detected automatically.", "Population data is read from the configured Google Sheet."])

    if population_long.empty: return

    ward_incidence, disease_ward_incidence, monthly_incidence, monthly_disease_incidence, yearly_incidence, audit = _build_incidence_bundle(case_df, population_long, denominator)

    # --- 2. POPULATION TREND ---
    st.divider()
    st.markdown("### 2. 👥 Population Trend & Matching Quality")
    
    population_tab, audit_tab = st.tabs(["📈 Population Trend", "🔍 Population Matching Audit"])
    with population_tab:
        wards = _alphabetical(population_long["WARD"].unique())
        selected_wards = _compact_multiselect("Select Ward(s)", wards, default=wards[:min(5, len(wards))], key="phase8b_pop_wards_state", defer=True)
        if selected_wards:
            chart_df = population_long[population_long["WARD"].isin(selected_wards)].copy().sort_values(["Population_Year", "WARD"])
            chart_df["Year"] = chart_df["Population_Year"].astype(int).astype(str)
            _render_line_chart(chart_df, "Year", "WARD", "Population", "Ward-wise Population Trend", "Population", height=450)
            with st.expander("📋 View Population Trend Data", expanded=False): st.dataframe(chart_df[["WARD", "Population_Year", "Population"]], use_container_width=True, hide_index=True)
        else: st.info("Select at least one ward.")
    with audit_tab:
        if not audit.empty:
            c1, c2, c3, c4 = st.columns(4)
            with c1: st.metric("Exact Matches", f"{(audit['Status'] == 'Exact year').sum():,}")
            with c2: st.metric("Previous-Year Fallback", f"{(audit['Status'] == 'Previous-year fallback').sum():,}")
            with c3: st.metric("Future-Year Fallback", f"{(audit['Status'] == 'Future-year fallback').sum():,}")
            with c4: st.metric("Unavailable", f"{(audit['Status'] == 'Population unavailable').sum():,}")
            st.dataframe(audit, use_container_width=True, hide_index=True)

    # --- 3. WARD-WISE INCIDENCE ---
    st.divider()
    st.markdown("### 3. 📍 Ward-wise Incidence Analysis")
    if not ward_incidence.empty:
        available_years = sorted(ward_incidence["Year"].unique())
        selected_year = _smart_year_slider("Select Analysis Year", available_years, default=available_years[-1], key="phase8b_ward_year_slider_state")
        if selected_year is not None:
            year_df = ward_incidence[ward_incidence["Year"] == selected_year].copy()
            ward_options = _alphabetical(year_df["Ward"].unique())
            selected_ward_analysis = _compact_multiselect("Select Ward(s)", ward_options, default=ward_options, key="phase8b_ward_sel_state", defer=True)
            if selected_ward_analysis:
                display_df = year_df[year_df["Ward"].isin(selected_ward_analysis)].copy().sort_values("Incidence Rate", ascending=False, na_position="last")
                _render_bar_chart(display_df, "Ward", "Incidence Rate", f"Ward-wise Incidence Rate — {selected_year}", f"Incidence per {denominator:,}", height=460)
                with st.expander("📋 View Ward-wise Incidence Data", expanded=False): st.dataframe(display_df[["Ward", "Year", "Cases", "Population", "Population Match Status", "Incidence Rate"]], use_container_width=True, hide_index=True)
                
                ward_disease = disease_ward_incidence[(disease_ward_incidence["Year"] == selected_year) & (disease_ward_incidence["Ward"].isin(selected_ward_analysis))].copy()
                if not ward_disease.empty:
                    ward_disease["Ward | Disease"] = ward_disease["Ward"] + " | " + ward_disease["Disease"]
                    top_ward_disease = ward_disease.sort_values("Incidence Rate", ascending=False).head(20)
                    _render_bar_chart(top_ward_disease, "Ward | Disease", "Incidence Rate", "Disease-specific Incidence within Selected Ward(s)", f"Incidence per {denominator:,}", height=500)

    # --- 4. DISEASE-WISE INCIDENCE ---
    st.divider()
    st.markdown("### 4. 🦠 Disease-wise Incidence Analysis")
    if not disease_ward_incidence.empty:
        disease_years = sorted(disease_ward_incidence["Year"].unique())
        disease_year = _smart_year_slider("Disease Analysis Year", disease_years, default=disease_years[-1], key="phase8b_disease_year_slider_state")
        if disease_year is not None:
            disease_source = disease_ward_incidence[disease_ward_incidence["Year"] == disease_year].copy()
            disease_options = disease_source["Disease"].value_counts().index.tolist()
            selected_diseases = _compact_multiselect("Select Disease(s)", disease_options, default=disease_options[:min(8, len(disease_options))], key="phase8b_disease_sel_state", defer=True)
            if selected_diseases:
                selected_source = disease_source[disease_source["Disease"].isin(selected_diseases)].copy()
                summary_rows = [{"Disease": d, "Cases": int(g["Cases"].sum()), "Population": float(g["Population"].sum()), "Incidence Rate": round(int(g["Cases"].sum())/float(g["Population"].sum())*denominator, 2)} for d, g in selected_source.groupby("Disease", observed=True) if not g[g["Population"].gt(0)].empty]
                disease_summary = pd.DataFrame(summary_rows).sort_values("Incidence Rate", ascending=False)
                if not disease_summary.empty:
                    _render_bar_chart(disease_summary, "Disease", "Incidence Rate", f"Disease-wise Incidence — {disease_year}", f"Incidence per {denominator:,}", height=450)
                    with st.expander("📋 View Disease-wise Data", expanded=False): st.dataframe(disease_summary, use_container_width=True, hide_index=True)
                
                disease_ward_options = _alphabetical(selected_source["Ward"].unique())
                selected_disease_wards = _compact_multiselect("Select Ward(s) for Disease Comparison", disease_ward_options, default=disease_ward_options[:min(8, len(disease_ward_options))], key="phase8b_dis_ward_sel_state", defer=True)
                if selected_disease_wards:
                    ward_disease_df = selected_source[selected_source["Ward"].isin(selected_disease_wards)].copy()
                    ward_disease_df["Ward | Disease"] = ward_disease_df["Ward"] + " | " + ward_disease_df["Disease"]
                    _render_bar_chart(ward_disease_df.sort_values("Incidence Rate", ascending=False), "Ward | Disease", "Incidence Rate", "Ward-wise Disease Incidence Comparison", f"Incidence per {denominator:,}", height=500)

    # --- 5. INCIDENCE TRENDS ---
    st.divider()
    st.markdown("### 5. 📈 Incidence Trends")
    trend_disease_options = ["All Diseases"] + (_alphabetical(monthly_disease_incidence["Disease"].unique()) if not monthly_disease_incidence.empty and "Disease" in monthly_disease_incidence.columns else [])
    selected_trend_disease = _single_select("Select Disease", trend_disease_options, default="All Diseases", key="phase8b_inc_trend_dis_state") or "All Diseases"

    monthly_tab, yearly_tab = st.tabs(["📅 Monthly Trend", "🗓️ Year-wise Trend"])
    with monthly_tab:
        monthly_source = monthly_incidence if selected_trend_disease == "All Diseases" else monthly_disease_incidence[monthly_disease_incidence["Disease"] == selected_trend_disease].copy()
        monthly_chart = _aggregate_monthly_incidence(monthly_source, denominator)
        if not monthly_chart.empty:
            monthly_chart["Series"] = selected_trend_disease
            _render_line_chart(monthly_chart, "Period", "Series", "Incidence Rate", f"Monthly Overall Incidence Trend — {selected_trend_disease}", f"Incidence per {denominator:,}", height=480)
        else: st.info("Monthly incidence data is not available.")
    with yearly_tab:
        chart_df = yearly_incidence.copy() if selected_trend_disease == "All Diseases" else _aggregate_year_incidence(disease_ward_incidence[disease_ward_incidence["Disease"] == selected_trend_disease].copy(), denominator)
        if not chart_df.empty:
            chart_df["Year Label"] = chart_df["Year"].astype(str)
            chart_df["Series"] = selected_trend_disease
            _render_line_chart(chart_df, "Year Label", "Series", "Incidence Rate", f"Year-wise Overall Incidence Trend — {selected_trend_disease}", f"Incidence per {denominator:,}", height=430)
        else: st.info("Year-wise incidence data is not available.")

    # --- 6. STATISTICAL SURVEILLANCE ---
    st.divider()
    st.markdown("### 6. 📐 Statistical Surveillance Analysis")
    
    baseline = _single_select("Statistical Baseline", ["All Available Years", "Last 2 Years", "Last 3 Years", "Last 5 Years"], default="All Available Years", key="phase8b_stat_baseline_state")
    statistical_source = _apply_baseline(case_df, baseline)

    disease_options = ["All Diseases"] + (_alphabetical(statistical_source[_valid_text(statistical_source["Disease"])]["Disease"].unique()) if "Disease" in statistical_source.columns else [])
    ward_options = ["All Wards"] + (_alphabetical(statistical_source[_valid_text(statistical_source["Ward Name"])]["Ward Name"].unique()) if "Ward Name" in statistical_source.columns else [])

    c1, c2 = st.columns(2)
    with c1: selected_stat_disease = _single_select("Select Disease", disease_options, default="All Diseases", key="phase8b_stat_disease_state")
    with c2: selected_stat_ward = _single_select("Select Ward", ward_options, default="All Wards", key="phase8b_stat_ward_state")

    if selected_stat_disease != "All Diseases" and "Disease" in statistical_source.columns: statistical_source = statistical_source[statistical_source["Disease"] == selected_stat_disease]
    if selected_stat_ward != "All Wards" and "Ward Name" in statistical_source.columns: statistical_source = statistical_source[statistical_source["Ward Name"] == selected_stat_ward]

    monthly_stats = _monthly_case_series(statistical_source)
    statistics = _threshold_statistics(monthly_stats)

    if statistics is not None and not monthly_stats.empty:
        c1, c2, c3, c4, c5 = st.columns(5)
        with c1: st.metric("Mean", f"{statistics['Mean']:,.2f}")
        with c2: st.metric("SD", f"{statistics['SD']:,.2f}")
        with c3: st.metric("Mean + 1 SD", f"{statistics['Mean + 1 SD']:,.2f}")
        with c4: st.metric("Mean + 2 SD", f"{statistics['Mean + 2 SD']:,.2f}")
        with c5: st.metric("Mean + 3 SD", f"{statistics['Mean + 3 SD']:,.2f}")

        title_suffix = ""
        if selected_stat_disease != "All Diseases": title_suffix += f" — {selected_stat_disease}"
        if selected_stat_ward != "All Wards": title_suffix += f" — Ward {selected_stat_ward}"
        _render_threshold_chart(monthly_stats, statistics, f"Monthly Cases{title_suffix} — Statistical Surveillance")
    else: st.info("Sufficient monthly data is not available for this statistical analysis.")

    # --- 7. DOWNLOAD REPORTS (Professional Excel, Word, PDF with CHARTS) ---
    st.divider()
    st.markdown("### 7. 📥 Download Reports & Data")
    st.caption("Generate a formatted professional report including data tables and all visible charts.")
    
    col1, col2 = st.columns(2)
    with col1: report_format = st.radio("Select Report Format:", ["Excel", "Word", "PDF"], horizontal=True)
    with col2: include_tables = st.radio("Include Data Tables?:", ["With Tables", "Without Tables (Summary Only)"], horizontal=True)
    
    # 1. EXCEL EXPORT
    if report_format == "Excel":
        output = io.BytesIO()
        try:
            with pd.ExcelWriter(output, engine="openpyxl") as writer:
                pd.DataFrame({"Metric": ["Total Records", "Wards Analyzed", "Denominator Used"], "Value": [len(case_df), population_wards, f"Per {denominator:,}"]}).to_excel(writer, sheet_name="Summary", index=False)
                if include_tables == "With Tables":
                    if not ward_incidence.empty: ward_incidence.to_excel(writer, sheet_name="Ward Incidence", index=False)
                    if not disease_ward_incidence.empty: disease_ward_incidence.to_excel(writer, sheet_name="Disease Incidence", index=False)
                    if not monthly_incidence.empty: monthly_incidence.to_excel(writer, sheet_name="Monthly Trend", index=False)
            st.download_button(label="📥 Download Excel Report", data=output.getvalue(), file_name="Incidence_Report.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        except Exception as e: st.error(f"Excel error: {e}")
            
    # 2. WORD EXPORT (Professional Layout)
    elif report_format == "Word":
        try:
            from docx import Document
            from docx.shared import Inches
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            doc = Document()
            title = doc.add_heading('Public Health Incidence & Surveillance Report', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_heading('Executive Summary', level=1)
            doc.add_paragraph(f"• Total Records Analyzed: {len(case_df):,}")
            doc.add_paragraph(f"• Denominator Used: Per {denominator:,} Population")
            
            if "export_charts" in st.session_state and st.session_state["export_charts"]:
                doc.add_heading('Visualizations', level=1)
                for chart_title, chart_obj in st.session_state["export_charts"].items():
                    doc.add_heading(chart_title, level=2)
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmp:
                        # Fixed width for HD resolution in Word
                        export_chart = chart_obj.properties(width=650)
                        export_chart.save(tmp.name, format="png", engine="vl-convert")
                        doc.add_picture(tmp.name, width=Inches(6.0)) # Perfect fit
            
            if include_tables == "With Tables" and not ward_incidence.empty:
                doc.add_heading('Ward Incidence Data (Top Records)', level=1)
                table = doc.add_table(rows=1, cols=3)
                table.style = 'Table Grid'
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text, hdr_cells[1].text, hdr_cells[2].text = 'Ward', 'Cases', 'Incidence Rate'
                for _, row in ward_incidence.head(40).iterrows(): 
                    row_cells = table.add_row().cells
                    row_cells[0].text, row_cells[1].text, row_cells[2].text = str(row['Ward']), str(row['Cases']), str(row['Incidence Rate'])
                
            output = io.BytesIO()
            doc.save(output)
            st.download_button(label="📥 Download Professional Word Report", data=output.getvalue(), file_name="Incidence_Report.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        except Exception as e: st.error(f"Word generation failed. Ensure python-docx and vl-convert-python are installed. Error: {e}")

    # 3. PDF EXPORT (Professional Layout)
    elif report_format == "PDF":
        try:
            from fpdf import FPDF
            def clean_text(text): return str(text).replace("—", "-").replace("–", "-").encode('latin-1', 'replace').decode('latin-1')

            class PDFReport(FPDF):
                def header(self):
                    self.set_fill_color(41, 128, 185) # BMC Blue
                    self.set_text_color(255, 255, 255)
                    self.set_font("Arial", 'B', 16)
                    self.cell(0, 15, " Public Health Incidence & Surveillance Report", ln=True, fill=True, align='C')
                    self.ln(5)
                def footer(self):
                    self.set_y(-15)
                    self.set_font("Arial", 'I', 8)
                    self.set_text_color(128, 128, 128)
                    self.cell(0, 10, f"Page {self.page_no()}", align='C')

            pdf = PDFReport()
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.add_page()
            pdf.set_font("Arial", 'B', 14)
            pdf.cell(0, 10, txt="Executive Summary", ln=True)
            pdf.set_font("Arial", size=12)
            pdf.cell(0, 8, txt=f"Total Records Analyzed: {len(case_df):,}", ln=True)
            pdf.cell(0, 8, txt=f"Denominator Used: Per {denominator:,} Population", ln=True)
            pdf.ln(5)
            
            if "export_charts" in st.session_state and st.session_state["export_charts"]:
                for chart_title, chart_obj in st.session_state["export_charts"].items():
                    pdf.set_font("Arial", 'B', 12)
                    pdf.cell(0, 10, txt=clean_text(chart_title), ln=True)
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmp:
                        # Force width to 700px for HD resolution
                        export_chart = chart_obj.properties(width=700)
                        export_chart.save(tmp.name, format="png", engine="vl-convert")
                        # 190w fits A4 perfectly with 10mm margins
                        pdf.image(tmp.name, x=10, w=190) 
                        pdf.ln(8)
            
            if include_tables == "With Tables" and not ward_incidence.empty:
                pdf.add_page()
                pdf.set_font("Arial", 'B', 14)
                pdf.cell(0, 10, txt="Data Table (Top Records)", ln=True)
                pdf.set_fill_color(200, 220, 255)
                pdf.set_font("Arial", 'B', 10)
                pdf.cell(90, 10, "Ward", border=1, fill=True)
                pdf.cell(40, 10, "Cases", border=1, align='C', fill=True)
                pdf.cell(60, 10, "Incidence Rate", border=1, align='C', fill=True)
                pdf.ln()
                pdf.set_font("Arial", size=10)
                for _, row in ward_incidence.head(40).iterrows(): 
                    pdf.cell(90, 10, clean_text(row['Ward'])[:45], border=1) # Limit ward name length
                    pdf.cell(40, 10, str(row['Cases']), border=1, align='C')
                    pdf.cell(60, 10, str(row['Incidence Rate']), border=1, align='C')
                    pdf.ln()
                    
            pdf_bytes = pdf.output(dest='S').encode('latin1')
            st.download_button(label="📥 Download Professional PDF", data=pdf_bytes, file_name="Incidence_Report.pdf", mime="application/pdf")
        except Exception as e: st.error(f"PDF generation failed. Error: {e}")

# ============================================================
# BACKWARD-COMPATIBLE ENTRY POINT
# ============================================================

def render_incidence(df):
    return render_incidence_analysis(df)
