from io import BytesIO
from datetime import datetime

import pandas as pd
import matplotlib.pyplot as plt

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


# ============================================================
# BRANDING
# ============================================================

TITLE = "MSU Mumbai Public Health Surveillance Dashboard"
SUBTITLE = "Surveillance • Monitoring • Analysis • Management"

REPORT_TITLE = "Dashboard Management Report"


# ============================================================
# COLORS
# ============================================================

NAVY = RGBColor(31, 56, 100)
BLUE = RGBColor(68, 114, 196)
LIGHT_BLUE = RGBColor(221, 235, 247)
GREY = RGBColor(89, 89, 89)
LIGHT_GREY = RGBColor(242, 242, 242)
DARK = RGBColor(35, 35, 35)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(70, 120, 70)
RED = RGBColor(170, 70, 70)


# ============================================================
# BASIC HELPERS
# ============================================================

def _safe_text(value):
    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass

    return str(value)


def _format_number(value):
    try:
        return f"{int(value):,}"
    except Exception:
        return _safe_text(value)


def _clean_series(series):
    if series is None:
        return pd.Series(dtype="string")

    return (
        series
        .fillna("")
        .astype(str)
        .str.strip()
    )


def _add_textbox(
    slide,
    text,
    left,
    top,
    width,
    height,
    font_size=18,
    bold=False,
    color=DARK,
    align=PP_ALIGN.LEFT,
):
    textbox = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )

    tf = textbox.text_frame
    tf.clear()
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.alignment = align

    run = p.add_run()
    run.text = _safe_text(text)
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Aptos"

    return textbox


def _add_title(slide, title, subtitle=None):
    _add_textbox(
        slide,
        title,
        0.55,
        0.25,
        12.2,
        0.45,
        font_size=23,
        bold=True,
        color=NAVY,
    )

    if subtitle:
        _add_textbox(
            slide,
            subtitle,
            0.58,
            0.72,
            12.0,
            0.30,
            font_size=10,
            color=GREY,
        )


def _add_footer(slide, slide_number):
    # footer line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.45),
        Inches(7.05),
        Inches(12.4),
        Inches(0.015),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT_BLUE
    line.line.fill.background()

    _add_textbox(
        slide,
        SUBTITLE,
        0.50,
        7.10,
        9.5,
        0.25,
        font_size=8,
        color=GREY,
    )

    _add_textbox(
        slide,
        f"Page {slide_number}",
        11.25,
        7.10,
        1.4,
        0.25,
        font_size=8,
        color=GREY,
        align=PP_ALIGN.RIGHT,
    )


def _add_section_box(
    slide,
    title,
    value,
    left,
    top,
    width,
    height=0.95,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_GREY
    shape.line.color.rgb = LIGHT_BLUE

    _add_textbox(
        slide,
        title,
        left + 0.12,
        top + 0.10,
        width - 0.24,
        0.25,
        font_size=9,
        bold=True,
        color=GREY,
    )

    _add_textbox(
        slide,
        value,
        left + 0.12,
        top + 0.36,
        width - 0.24,
        0.42,
        font_size=19,
        bold=True,
        color=NAVY,
    )

    return shape


# ============================================================
# TABLE
# ============================================================

def _add_table(
    slide,
    dataframe,
    left,
    top,
    width,
    height,
    font_size=9,
):
    if dataframe is None or dataframe.empty:
        _add_textbox(
            slide,
            "No data available for the selected filters.",
            left,
            top,
            width,
            0.5,
            font_size=12,
            color=GREY,
        )
        return

    data = dataframe.copy()

    rows = len(data) + 1
    cols = len(data.columns)

    table_shape = slide.shapes.add_table(
        rows,
        cols,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )

    table = table_shape.table

    for i, column in enumerate(data.columns):
        cell = table.cell(0, i)
        cell.text = _safe_text(column)

        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY

        for paragraph in cell.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER

            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(font_size)
                run.font.color.rgb = WHITE
                run.font.name = "Aptos"

    for r in range(len(data)):
        for c in range(cols):
            value = data.iloc[r, c]

            if isinstance(value, float):
                if pd.isna(value):
                    text = ""
                else:
                    text = f"{value:,.2f}"
            else:
                text = _safe_text(value)

            cell = table.cell(r + 1, c)
            cell.text = text

            if r % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_GREY

            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER

                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = DARK
                    run.font.name = "Aptos"

    # equal-ish widths
    for c in range(cols):
        table.columns[c].width = Inches(width / max(cols, 1))

    return table_shape


# ============================================================
# CHART
# ============================================================

def _create_bar_chart(dataframe, category_col, value_col, title):
    if dataframe is None or dataframe.empty:
        return None

    data = dataframe.copy()

    if category_col not in data.columns:
        return None

    if value_col not in data.columns:
        return None

    data[value_col] = pd.to_numeric(
        data[value_col],
        errors="coerce",
    )

    data = data.dropna(subset=[value_col])

    if data.empty:
        return None

    data = data.head(10)

    fig, ax = plt.subplots(figsize=(8.5, 4.2))

    ax.bar(
        data[category_col].astype(str),
        data[value_col],
    )

    ax.set_title(
        title,
        fontsize=13,
        fontweight="bold",
    )

    ax.set_ylabel("Records")

    ax.tick_params(
        axis="x",
        rotation=35,
        labelsize=8,
    )

    ax.grid(
        axis="y",
        alpha=0.20,
    )

    fig.tight_layout()

    buffer = BytesIO()

    fig.savefig(
        buffer,
        format="png",
        dpi=180,
        bbox_inches="tight",
    )

    plt.close(fig)

    buffer.seek(0)

    return buffer


def _create_line_chart(dataframe, x_col, y_col, title):
    if dataframe is None or dataframe.empty:
        return None

    if x_col not in dataframe.columns:
        return None

    if y_col not in dataframe.columns:
        return None

    data = dataframe.copy()

    data[y_col] = pd.to_numeric(
        data[y_col],
        errors="coerce",
    )

    data = data.dropna(subset=[y_col])

    if data.empty:
        return None

    fig, ax = plt.subplots(figsize=(8.5, 4.2))

    ax.plot(
        data[x_col].astype(str),
        data[y_col],
        marker="o",
        linewidth=2,
    )

    ax.set_title(
        title,
        fontsize=13,
        fontweight="bold",
    )

    ax.set_ylabel("Records")

    ax.tick_params(
        axis="x",
        rotation=35,
        labelsize=8,
    )

    ax.grid(
        alpha=0.20,
    )

    fig.tight_layout()

    buffer = BytesIO()

    fig.savefig(
        buffer,
        format="png",
        dpi=180,
        bbox_inches="tight",
    )

    plt.close(fig)

    buffer.seek(0)

    return buffer


# ============================================================
# TITLE SLIDE
# ============================================================

def _add_cover_slide(
    prs,
    df,
    report_period=None,
    filter_summary=None,
):
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    # top band
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(13.333),
        Inches(1.25),
    )

    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()

    _add_textbox(
        slide,
        TITLE,
        0.75,
        1.75,
        11.8,
        0.75,
        font_size=28,
        bold=True,
        color=NAVY,
        align=PP_ALIGN.CENTER,
    )

    _add_textbox(
        slide,
        REPORT_TITLE,
        1.0,
        2.65,
        11.3,
        0.55,
        font_size=22,
        bold=True,
        color=BLUE,
        align=PP_ALIGN.CENTER,
    )

    _add_textbox(
        slide,
        SUBTITLE,
        1.0,
        3.30,
        11.3,
        0.45,
        font_size=14,
        color=GREY,
        align=PP_ALIGN.CENTER,
    )

    period_text = (
        f"Reporting Period: {report_period}"
        if report_period
        else "Reporting Period: As per selected dashboard filters"
    )

    _add_textbox(
        slide,
        period_text,
        1.0,
        4.05,
        11.3,
        0.40,
        font_size=13,
        color=DARK,
        align=PP_ALIGN.CENTER,
    )

    _add_textbox(
        slide,
        f"Records included in report: {_format_number(len(df))}",
        1.0,
        4.55,
        11.3,
        0.40,
        font_size=13,
        color=DARK,
        align=PP_ALIGN.CENTER,
    )

    if filter_summary:
        _add_textbox(
            slide,
            f"Filter scope: {filter_summary}",
            1.0,
            5.05,
            11.3,
            0.75,
            font_size=10,
            color=GREY,
            align=PP_ALIGN.CENTER,
        )

    _add_textbox(
        slide,
        f"Generated: {datetime.now().strftime('%d-%m-%Y %H:%M')}",
        1.0,
        6.25,
        11.3,
        0.35,
        font_size=10,
        color=GREY,
        align=PP_ALIGN.CENTER,
    )


# ============================================================
# EXECUTIVE SUMMARY
# ============================================================

def _add_summary_slide(
    prs,
    df,
    report_period=None,
):
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    _add_title(
        slide,
        "1. Executive Summary",
        "Summary of the currently selected dashboard data",
    )

    total_records = len(df)

    diseases = (
        df["Disease"].replace("", pd.NA).dropna().nunique()
        if "Disease" in df.columns
        else 0
    )

    facilities = (
        df["Facility Name"].replace("", pd.NA).dropna().nunique()
        if "Facility Name" in df.columns
        else 0
    )

    wards = (
        df["Ward Name"].replace("", pd.NA).dropna().nunique()
        if "Ward Name" in df.columns
        else 0
    )

    _add_section_box(
        slide,
        "Total records",
        _format_number(total_records),
        0.65,
        1.25,
        2.85,
    )

    _add_section_box(
        slide,
        "Diseases",
        _format_number(diseases),
        3.65,
        1.25,
        2.85,
    )

    _add_section_box(
        slide,
        "Facilities",
        _format_number(facilities),
        6.65,
        1.25,
        2.85,
    )

    _add_section_box(
        slide,
        "Wards",
        _format_number(wards),
        9.65,
        1.25,
        2.85,
    )

    if "Disease" in df.columns and not df.empty:
        disease_counts = (
            _clean_series(df["Disease"])
            .replace("", pd.NA)
            .dropna()
            .value_counts()
            .head(10)
            .reset_index()
        )

        disease_counts.columns = [
            "Disease",
            "Records",
        ]

        _add_textbox(
            slide,
            "Disease-wise burden",
            0.65,
            2.55,
            5.8,
            0.35,
            font_size=15,
            bold=True,
            color=NAVY,
        )

        _add_table(
            slide,
            disease_counts,
            0.65,
            2.95,
            5.8,
            2.85,
            font_size=8,
        )

    if "Facility Name" in df.columns and not df.empty:
        facility_counts = (
            _clean_series(df["Facility Name"])
            .replace("", pd.NA)
            .dropna()
            .value_counts()
            .head(10)
            .reset_index()
        )

        facility_counts.columns = [
            "Facility",
            "Records",
        ]

        _add_textbox(
            slide,
            "Facility-wise burden",
            6.75,
            2.55,
            5.8,
            0.35,
            font_size=15,
            bold=True,
            color=NAVY,
        )

        _add_table(
            slide,
            facility_counts,
            6.75,
            2.95,
            5.8,
            2.85,
            font_size=8,
        )

    _add_textbox(
        slide,
        (
            "Management note: The figures in this presentation represent "
            "the records remaining after application of the dashboard's "
            "current Global Dashboard Control filters."
        ),
        0.75,
        6.05,
        11.8,
        0.65,
        font_size=10,
        color=GREY,
    )

    _add_footer(
        slide,
        len(prs.slides),
    )


# ============================================================
# DISEASE SLIDE
# ============================================================

def _add_disease_slide(prs, df):
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    _add_title(
        slide,
        "2. Disease-wise Analysis",
        "Distribution of records by disease / diagnosis",
    )

    if "Disease" not in df.columns:
        _add_textbox(
            slide,
            "Disease information is not available.",
            0.75,
            1.5,
            11,
            0.5,
            font_size=14,
            color=GREY,
        )
        _add_footer(slide, len(prs.slides))
        return

    counts = (
        _clean_series(df["Disease"])
        .replace("", pd.NA)
        .dropna()
        .value_counts()
        .reset_index()
    )

    counts.columns = [
        "Disease",
        "Records",
    ]

    if counts.empty:
        _add_textbox(
            slide,
            "No disease records available for the selected filters.",
            0.75,
            1.5,
            11,
            0.5,
            font_size=14,
            color=GREY,
        )
    else:
        chart = _create_bar_chart(
            counts,
            "Disease",
            "Records",
            "Top disease burden",
        )

        if chart:
            slide.shapes.add_picture(
                chart,
                Inches(0.65),
                Inches(1.35),
                width=Inches(7.3),
            )

        _add_table(
            slide,
            counts.head(10),
            8.15,
            1.45,
            4.45,
            4.75,
            font_size=8,
        )

    _add_footer(slide, len(prs.slides))


# ============================================================
# FACILITY SLIDE
# ============================================================

def _add_facility_slide(prs, df):
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    _add_title(
        slide,
        "3. Facility-wise Analysis",
        "Distribution of records across reporting facilities",
    )

    if "Facility Name" not in df.columns:
        _add_textbox(
            slide,
            "Facility information is not available.",
            0.75,
            1.5,
            11,
            0.5,
            font_size=14,
            color=GREY,
        )
        _add_footer(slide, len(prs.slides))
        return

    counts = (
        _clean_series(df["Facility Name"])
        .replace("", pd.NA)
        .dropna()
        .value_counts()
        .reset_index()
    )

    counts.columns = [
        "Facility",
        "Records",
    ]

    if counts.empty:
        _add_textbox(
            slide,
            "No facility records available.",
            0.75,
            1.5,
            11,
            0.5,
            font_size=14,
            color=GREY,
        )
    else:
        chart = _create_bar_chart(
            counts,
            "Facility",
            "Records",
            "Top facilities by record count",
        )

        if chart:
            slide.shapes.add_picture(
                chart,
                Inches(0.65),
                Inches(1.35),
                width=Inches(7.3),
            )

        _add_table(
            slide,
            counts.head(10),
            8.15,
            1.45,
            4.45,
            4.75,
            font_size=8,
        )

    _add_footer(slide, len(prs.slides))


# ============================================================
# WARD SLIDE
# ============================================================

def _add_ward_slide(prs, df):
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    _add_title(
        slide,
        "4. Ward-wise Analysis",
        "Distribution of records across wards",
    )

    if "Ward Name" not in df.columns:
        _add_textbox(
            slide,
            "Ward information is not available.",
            0.75,
            1.5,
            11,
            0.5,
            font_size=14,
            color=GREY,
        )
        _add_footer(slide, len(prs.slides))
        return

    counts = (
        _clean_series(df["Ward Name"])
        .replace("", pd.NA)
        .dropna()
        .value_counts()
        .reset_index()
    )

    counts.columns = [
        "Ward",
        "Records",
    ]

    if counts.empty:
        _add_textbox(
            slide,
            "No ward records available.",
            0.75,
            1.5,
            11,
            0.5,
            font_size=14,
            color=GREY,
        )
    else:
        chart = _create_bar_chart(
            counts,
            "Ward",
            "Records",
            "Top wards by record count",
        )

        if chart:
            slide.shapes.add_picture(
                chart,
                Inches(0.65),
                Inches(1.35),
                width=Inches(7.3),
            )

        _add_table(
            slide,
            counts.head(10),
            8.15,
            1.45,
            4.45,
            4.75,
            font_size=8,
        )

    _add_footer(slide, len(prs.slides))


# ============================================================
# DEMOGRAPHICS SLIDE
# ============================================================

def _add_demographics_slide(prs, df):
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    _add_title(
        slide,
        "5. Demographic Analysis",
        "Age and gender distribution within the selected data",
    )

    # AGE
    if "Age Group" in df.columns:
        age_counts = (
            _clean_series(df["Age Group"])
            .replace("", pd.NA)
            .dropna()
            .value_counts()
            .reset_index()
        )

        age_counts.columns = [
            "Age Group",
            "Records",
        ]

        _add_textbox(
            slide,
            "Age group distribution",
            0.65,
            1.25,
            5.7,
            0.35,
            font_size=15,
            bold=True,
            color=NAVY,
        )

        _add_table(
            slide,
            age_counts,
            0.65,
            1.70,
            5.7,
            4.65,
            font_size=8,
        )

    # GENDER
    if "Gender" in df.columns:
        gender_counts = (
            _clean_series(df["Gender"])
            .replace("", pd.NA)
            .dropna()
            .value_counts()
            .reset_index()
        )

        gender_counts.columns = [
            "Gender",
            "Records",
        ]

        _add_textbox(
            slide,
            "Gender distribution",
            6.75,
            1.25,
            5.7,
            0.35,
            font_size=15,
            bold=True,
            color=NAVY,
        )

        _add_table(
            slide,
            gender_counts,
            6.75,
            1.70,
            5.7,
            4.65,
            font_size=8,
        )

    _add_footer(slide, len(prs.slides))


# ============================================================
# MONTHLY TREND
# ============================================================

def _add_monthly_slide(prs, df):
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    _add_title(
        slide,
        "6. Monthly Analysis and Trend",
        "Month-wise distribution of selected records",
    )

    if "Month" not in df.columns:
        _add_textbox(
            slide,
            "Monthly information is not available.",
            0.75,
            1.5,
            11,
            0.5,
            font_size=14,
            color=GREY,
        )
        _add_footer(slide, len(prs.slides))
        return

    month_counts = (
        _clean_series(df["Month"])
        .replace("", pd.NA)
        .dropna()
        .value_counts()
        .reset_index()
    )

    month_counts.columns = [
        "Month",
        "Records",
    ]

    # If Year exists, use chronological Year-Month
    if "Year" in df.columns:
        monthly = (
            df.assign(
                _year=pd.to_numeric(
                    df["Year"],
                    errors="coerce",
                )
            )
            .groupby(
                ["_year", "Month"],
                dropna=False,
            )
            .size()
            .reset_index(
                name="Records"
            )
        )

        monthly = monthly[
            monthly["_year"].notna()
        ].copy()

        if not monthly.empty:
            monthly["_label"] = (
                monthly["_year"]
                .astype(int)
                .astype(str)
                + " - "
                + monthly["Month"].astype(str)
            )

            monthly = monthly[
                ["_label", "Records"]
            ]

            monthly.columns = [
                "Month",
                "Records",
            ]

            month_counts = monthly

    chart = _create_line_chart(
        month_counts,
        "Month",
        "Records",
        "Monthly record trend",
    )

    if chart:
        slide.shapes.add_picture(
            chart,
            Inches(0.65),
            Inches(1.35),
            width=Inches(7.3),
        )

    _add_table(
        slide,
        month_counts,
        8.15,
        1.45,
        4.45,
        4.75,
        font_size=8,
    )

    _add_footer(slide, len(prs.slides))


# ============================================================
# OPD/IPD SLIDE
# ============================================================

def _add_opdipd_slide(prs, df):
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    _add_title(
        slide,
        "7. OPD / IPD Analysis",
        "Distribution by service / admission category",
    )

    column = None

    if "OPD/IPD" in df.columns:
        column = "OPD/IPD"
    elif "Opd Ipd" in df.columns:
        column = "Opd Ipd"

    if column is None:
        _add_textbox(
            slide,
            "OPD/IPD information is not available.",
            0.75,
            1.5,
            11,
            0.5,
            font_size=14,
            color=GREY,
        )
        _add_footer(slide, len(prs.slides))
        return

    counts = (
        _clean_series(df[column])
        .replace("", pd.NA)
        .dropna()
        .value_counts()
        .reset_index()
    )

    counts.columns = [
        "Category",
        "Records",
    ]

    chart = _create_bar_chart(
        counts,
        "Category",
        "Records",
        "OPD / IPD distribution",
    )

    if chart:
        slide.shapes.add_picture(
            chart,
            Inches(0.65),
            Inches(1.35),
            width=Inches(7.3),
        )

    _add_table(
        slide,
        counts,
        8.15,
        1.45,
        4.45,
        4.75,
        font_size=8,
    )

    _add_footer(slide, len(prs.slides))


# ============================================================
# DATA QUALITY SLIDE
# ============================================================

def _add_quality_slide(prs, df):
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    _add_title(
        slide,
        "8. Data Quality and Validation",
        "Basic completeness indicators for the selected dataset",
    )

    columns_to_check = [
        "Reporting Date",
        "Disease",
        "Facility Name",
        "Ward Name",
        "Gender",
        "Age",
        "Age Group",
        "OPD/IPD",
    ]

    rows = []

    total = len(df)

    for column in columns_to_check:
        if column not in df.columns:
            continue

        missing = df[column].isna().sum()

        if df[column].dtype == "object":
            missing += (
                df[column]
                .astype(str)
                .str.strip()
                .eq("")
                .sum()
            )

        completeness = (
            ((total - missing) / total * 100)
            if total > 0
            else 0
        )

        rows.append(
            {
                "Field": column,
                "Complete Records": total - missing,
                "Completeness": f"{completeness:.1f}%",
            }
        )

    quality_df = pd.DataFrame(rows)

    if quality_df.empty:
        _add_textbox(
            slide,
            "No validation fields available.",
            0.75,
            1.5,
            11,
            0.5,
            font_size=14,
            color=GREY,
        )
    else:
        _add_table(
            slide,
            quality_df,
            0.75,
            1.35,
            11.8,
            4.85,
            font_size=9,
        )

    _add_textbox(
        slide,
        (
            "Interpretation: Completeness percentages indicate the proportion "
            "of selected records containing a usable value in each field. "
            "They should be interpreted together with the detailed "
            "Validation & KPI dashboard."
        ),
        0.75,
        6.35,
        11.8,
        0.55,
        font_size=9,
        color=GREY,
    )

    _add_footer(slide, len(prs.slides))


# ============================================================
# MANAGEMENT FINDINGS
# ============================================================

def _add_findings_slide(prs, df):
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    _add_title(
        slide,
        "9. Key Management Findings",
        "Data-driven observations from the currently selected records",
    )

    findings = []

    if "Disease" in df.columns:
        disease = (
            _clean_series(df["Disease"])
            .replace("", pd.NA)
            .dropna()
            .value_counts()
        )

        if not disease.empty:
            findings.append(
                f"Highest recorded disease category: "
                f"{disease.index[0]} "
                f"({_format_number(disease.iloc[0])} records)."
            )

    if "Facility Name" in df.columns:
        facility = (
            _clean_series(df["Facility Name"])
            .replace("", pd.NA)
            .dropna()
            .value_counts()
        )

        if not facility.empty:
            findings.append(
                f"Highest record volume by facility: "
                f"{facility.index[0]} "
                f"({_format_number(facility.iloc[0])} records)."
            )

    if "Ward Name" in df.columns:
        ward = (
            _clean_series(df["Ward Name"])
            .replace("", pd.NA)
            .dropna()
            .value_counts()
        )

        if not ward.empty:
            findings.append(
                f"Highest record volume by ward: "
                f"{ward.index[0]} "
                f"({_format_number(ward.iloc[0])} records)."
            )

    if "Gender" in df.columns:
        gender = (
            _clean_series(df["Gender"])
            .replace("", pd.NA)
            .dropna()
            .value_counts()
        )

        if not gender.empty:
            findings.append(
                f"Most represented gender category in the selected records: "
                f"{gender.index[0]} "
                f"({_format_number(gender.iloc[0])} records)."
            )

    if "Age Group" in df.columns:
        age = (
            _clean_series(df["Age Group"])
            .replace("", pd.NA)
            .dropna()
            .value_counts()
        )

        if not age.empty:
            findings.append(
                f"Most represented age group: "
                f"{age.index[0]} "
                f"({_format_number(age.iloc[0])} records)."
            )

    if not findings:
        findings.append(
            "No management findings could be generated from the selected data."
        )

    top = 0.0

    for i, finding in enumerate(findings):
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.85),
            Inches(1.35 + top),
            Inches(11.65),
            Inches(0.72),
        )

        box.fill.solid()
        box.fill.fore_color.rgb = LIGHT_GREY
        box.line.color.rgb = LIGHT_BLUE

        _add_textbox(
            slide,
            f"{i + 1}. {finding}",
            1.05,
            1.53 + top,
            11.25,
            0.38,
            font_size=12,
            color=DARK,
        )

        top += 0.88

        if top > 4.8:
            break

    _add_textbox(
        slide,
        (
            "These observations are descriptive summaries of the selected "
            "dashboard data and should be considered along with programme "
            "context, field verification and the detailed dashboard views."
        ),
        0.85,
        6.15,
        11.65,
        0.55,
        font_size=9,
        color=GREY,
    )

    _add_footer(slide, len(prs.slides))


# ============================================================
# CONCLUSION
# ============================================================

def _add_conclusion_slide(prs, df):
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    _add_title(
        slide,
        "10. Conclusion and Management Use",
        "Summary of the report scope",
    )

    points = [
        "This presentation is generated directly from the records selected through the dashboard's Global Dashboard Control.",
        "The report reflects the current filtered dataset and does not apply an independent filter.",
        "Facility-wise, ward-wise, demographic and temporal patterns can be reviewed together for programme monitoring.",
        "Data quality indicators should be reviewed before using the findings for operational decisions.",
        "The dashboard remains the primary interactive source for further drill-down, validation and record-level review.",
    ]

    for i, point in enumerate(points):
        _add_textbox(
            slide,
            f"• {point}",
            0.95,
            1.45 + (i * 0.82),
            11.3,
            0.60,
            font_size=13,
            color=DARK,
        )

    _add_textbox(
        slide,
        (
            f"Total records included in this report: "
            f"{_format_number(len(df))}"
        ),
        0.95,
        5.85,
        11.3,
        0.40,
        font_size=12,
        bold=True,
        color=NAVY,
    )

    _add_textbox(
        slide,
        TITLE,
        0.95,
        6.35,
        11.3,
        0.35,
        font_size=10,
        color=GREY,
    )

    _add_footer(
        slide,
        len(prs.slides),
    )


# ============================================================
# MAIN PPT GENERATOR
# ============================================================

def generate_ppt_report(
    df,
    report_period=None,
    filter_summary=None,
):
    """
    Generate a PowerPoint management report from the
    currently filtered dashboard dataframe.

    Parameters
    ----------
    df : pandas.DataFrame
        The SAME filtered dataframe used by the dashboard.

    report_period : str, optional
        Reporting period displayed in the PPT.

    filter_summary : str, optional
        Human-readable description of current filters.

    Returns
    -------
    bytes
        PowerPoint file content.
    """

    if df is None:
        df = pd.DataFrame()

    df = df.copy()

    prs = Presentation()

    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # --------------------------------------------------------
    # Cover
    # --------------------------------------------------------

    _add_cover_slide(
        prs,
        df,
        report_period=report_period,
        filter_summary=filter_summary,
    )

    # --------------------------------------------------------
    # Main report sections
    # --------------------------------------------------------

    _add_summary_slide(
        prs,
        df,
        report_period=report_period,
    )

    _add_disease_slide(
        prs,
        df,
    )

    _add_facility_slide(
        prs,
        df,
    )

    _add_ward_slide(
        prs,
        df,
    )

    _add_demographics_slide(
        prs,
        df,
    )

    _add_monthly_slide(
        prs,
        df,
    )

    _add_opdipd_slide(
        prs,
        df,
    )

    _add_quality_slide(
        prs,
        df,
    )

    _add_findings_slide(
        prs,
        df,
    )

    _add_conclusion_slide(
        prs,
        df,
    )

    # --------------------------------------------------------
    # Save to memory
    # --------------------------------------------------------

    output = BytesIO()

    prs.save(output)

    output.seek(0)

    return output.getvalue()
